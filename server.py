"""FastAPI API for the personal Pomodoro app. PostgreSQL is the only record store."""
from __future__ import annotations

import csv
import io
import os
import re
import uuid
from datetime import date, datetime, timezone
from functools import lru_cache
from typing import Annotated, Any

import jwt
import psycopg
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from jwt import PyJWKClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from dotenv import load_dotenv

APP_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(APP_DIR, ".env"))
app = FastAPI(title="Pomodoro Matías Henquín", version="1.0.0")


class SessionInput(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    id_estudio: str | None = Field(default=None, max_length=80)
    fecha: date
    materia: str = Field(min_length=1, max_length=120)
    cantidad: int = Field(gt=0, le=100)
    tiempo: int = Field(gt=0, le=600)

    @field_validator("id_estudio")
    @classmethod
    def clean_id(cls, value: str | None) -> str | None:
        if value is not None and not value.strip():
            return None
        return value

    @field_validator("materia")
    @classmethod
    def clean_subject(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("La materia es obligatoria.")
        return value

    @model_validator(mode="after")
    def check_total(self) -> "SessionInput":
        if self.cantidad * self.tiempo > 60000:
            raise ValueError("La sesión excede el máximo permitido.")
        return self


def config() -> dict[str, str | None]:
    return {
        "supabaseUrl": os.getenv("SUPABASE_URL"),
        "supabaseAnonKey": os.getenv("SUPABASE_ANON_KEY"),
    }


def require_database() -> str:
    dsn = os.getenv("DATABASE_URL", "").strip()
    if not dsn:
        raise HTTPException(503, "La base PostgreSQL todavía no está configurada.")
    return dsn


def db_connection():
    try:
        return psycopg.connect(require_database(), connect_timeout=8, application_name="pomodoro-matias",
                               prepare_threshold=None, sslmode="require")
    except HTTPException:
        raise
    except psycopg.Error as exc:
        raise HTTPException(503, "No se pudo conectar con la base de datos de Pomodoro.") from exc


@lru_cache(maxsize=8)
def jwks_client(url: str) -> PyJWKClient:
    return PyJWKClient(url, cache_keys=True, lifespan=600, timeout=5)


def current_user(authorization: Annotated[str | None, Header()] = None) -> str:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Iniciá sesión para continuar.", headers={"WWW-Authenticate": "Bearer"})
    supabase_url = os.getenv("SUPABASE_URL", "").rstrip("/")
    allowed_user_id = os.getenv("POMODORO_ALLOWED_USER_ID", "").strip()
    if not supabase_url or not allowed_user_id:
        raise HTTPException(503, "Falta configurar el acceso personal de Pomodoro.")
    token = authorization.split(" ", 1)[1].strip()
    try:
        issuer = f"{supabase_url}/auth/v1"
        key = jwks_client(f"{issuer}/.well-known/jwks.json").get_signing_key_from_jwt(token).key
        claims = jwt.decode(
            token,
            key,
            algorithms=["ES256", "RS256", "EdDSA"],
            audience="authenticated",
            issuer=issuer,
            options={"require": ["exp", "sub", "iss", "aud"]},
            leeway=30,
        )
        user_id = str(uuid.UUID(claims["sub"]))
        expected_id = str(uuid.UUID(allowed_user_id))
        if user_id != expected_id:
            raise HTTPException(403, "Esta cuenta no está autorizada para Pomodoro.")
        return user_id
    except HTTPException:
        raise
    except (jwt.PyJWTError, ValueError, KeyError) as exc:
        raise HTTPException(401, "La sesión no es válida o venció.", headers={"WWW-Authenticate": "Bearer"}) from exc
    except Exception as exc:
        raise HTTPException(503, "No se pudo verificar la sesión con Supabase Auth.") from exc


UserId = Annotated[str, Depends(current_user)]


def row_to_dict(row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        "id_estudio": row[0], "fecha": row[1].isoformat(), "materia": row[2],
        "cantidad": row[3], "tiempo": row[4], "total_minutos": row[5],
        "created_at": row[6].isoformat() if row[6] else None,
        "updated_at": row[7].isoformat() if row[7] else None,
    }


def get_sessions(user_id: str) -> list[dict[str, Any]]:
    with db_connection() as connection:
        rows = connection.execute(
            """SELECT id_estudio, fecha, materia, cantidad, tiempo, total_minutos, created_at, updated_at
               FROM public.pomodoro_sessions WHERE user_id = %s ORDER BY fecha DESC, created_at DESC, id_estudio""",
            (user_id,),
        ).fetchall()
    return [row_to_dict(row) for row in rows]


def write_session(data: SessionInput, user_id: str, *, update_id: str | None = None) -> dict[str, Any]:
    record_id = update_id or data.id_estudio or str(uuid.uuid4())
    now = datetime.now(timezone.utc)
    values = (record_id, user_id, data.fecha, data.materia, data.cantidad, data.tiempo, now)
    with db_connection() as connection:
        if update_id:
            row = connection.execute(
                """UPDATE public.pomodoro_sessions SET fecha=%s,materia=%s,cantidad=%s,tiempo=%s,
                   updated_at=%s WHERE id_estudio=%s AND user_id=%s
                   RETURNING id_estudio,fecha,materia,cantidad,tiempo,total_minutos,created_at,updated_at""",
                (data.fecha, data.materia, data.cantidad, data.tiempo, now, update_id, user_id),
            ).fetchone()
            if not row:
                raise HTTPException(404, "No se encontró esa sesión.")
        else:
            row = connection.execute(
                """INSERT INTO public.pomodoro_sessions
                   (id_estudio,user_id,fecha,materia,cantidad,tiempo,created_at,updated_at)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                   ON CONFLICT (id_estudio) DO UPDATE SET fecha=EXCLUDED.fecha,materia=EXCLUDED.materia,
                     cantidad=EXCLUDED.cantidad,tiempo=EXCLUDED.tiempo,
                     updated_at=EXCLUDED.updated_at
                   WHERE public.pomodoro_sessions.user_id=EXCLUDED.user_id
                   RETURNING id_estudio,fecha,materia,cantidad,tiempo,total_minutos,created_at,updated_at""",
                values + (now,),
            ).fetchone()
            if not row:
                raise HTTPException(409, "Ese identificador ya pertenece a otra cuenta.")
    return row_to_dict(row)


@app.get("/api/config")
def public_config():
    """Only publishable Supabase client settings; never return database or service credentials."""
    return config()


@app.get("/api/health")
def health():
    return {"ok": True, "database_configured": bool(os.getenv("DATABASE_URL")), "auth_configured": bool(os.getenv("SUPABASE_URL") and os.getenv("SUPABASE_ANON_KEY") and os.getenv("POMODORO_ALLOWED_USER_ID"))}


@app.get("/api/sessions")
def list_sessions(user_id: UserId):
    return {"sessions": get_sessions(user_id)}


@app.post("/api/sessions", status_code=201)
def create_session(data: SessionInput, user_id: UserId):
    record = write_session(data, user_id)
    return {"session": record}


@app.put("/api/sessions/{session_id}")
def update_session(session_id: str, data: SessionInput, user_id: UserId):
    return {"session": write_session(data, user_id, update_id=session_id)}


@app.delete("/api/sessions/{session_id}", status_code=204)
def delete_session(session_id: str, user_id: UserId):
    with db_connection() as connection:
        result = connection.execute(
            "DELETE FROM public.pomodoro_sessions WHERE id_estudio=%s AND user_id=%s", (session_id, user_id)
        )
        if result.rowcount == 0:
            raise HTTPException(404, "No se encontró esa sesión.")
    return JSONResponse(status_code=204, content=None)


@app.get("/api/subjects")
def list_subjects(user_id: UserId):
    with db_connection() as connection:
        rows = connection.execute(
            "SELECT DISTINCT materia FROM public.pomodoro_sessions WHERE user_id=%s ORDER BY materia", (user_id,)
        ).fetchall()
    return {"subjects": [row[0] for row in rows]}


@app.get("/api/statistics")
def statistics(user_id: UserId):
    with db_connection() as connection:
        totals = connection.execute(
            """SELECT count(*), coalesce(sum(total_minutos),0), coalesce(sum(cantidad),0),
                      count(distinct materia), min(fecha), max(fecha)
               FROM public.pomodoro_sessions WHERE user_id=%s""", (user_id,)
        ).fetchone()
        subjects = connection.execute(
            """SELECT materia, sum(total_minutos)::integer AS total_minutos
               FROM public.pomodoro_sessions WHERE user_id=%s
               GROUP BY materia ORDER BY total_minutos DESC, materia""", (user_id,)
        ).fetchall()
        days = connection.execute(
            """SELECT fecha, sum(total_minutos)::integer AS total_minutos
               FROM public.pomodoro_sessions WHERE user_id=%s GROUP BY fecha ORDER BY fecha""", (user_id,)
        ).fetchall()
        weeks = connection.execute(
            """SELECT date_trunc('week', fecha)::date AS semana, sum(total_minutos)::integer AS total_minutos
               FROM public.pomodoro_sessions WHERE user_id=%s GROUP BY 1 ORDER BY 1""", (user_id,)
        ).fetchall()
        months = connection.execute(
            """SELECT date_trunc('month', fecha)::date AS mes, sum(total_minutos)::integer AS total_minutos
               FROM public.pomodoro_sessions WHERE user_id=%s GROUP BY 1 ORDER BY 1""", (user_id,)
        ).fetchall()
    return {
        "session_count": totals[0], "total_minutos": totals[1], "ciclos": totals[2],
        "materias": totals[3], "primera_fecha": totals[4].isoformat() if totals[4] else None,
        "ultima_fecha": totals[5].isoformat() if totals[5] else None,
        "por_materia": [{"materia": row[0], "total_minutos": row[1]} for row in subjects],
        "por_dia": [{"fecha": row[0].isoformat(), "total_minutos": row[1]} for row in days],
        "por_semana": [{"fecha": row[0].isoformat(), "total_minutos": row[1]} for row in weeks],
        "por_mes": [{"fecha": row[0].isoformat(), "total_minutos": row[1]} for row in months],
    }


def parse_minutes(value: str) -> int:
    value = value.strip()
    if re.fullmatch(r"\d{1,3}:\d{2}:\d{2}", value):
        hours, minutes, seconds = map(int, value.split(":"))
        if minutes > 59 or seconds > 59:
            raise ValueError("Duración horaria inválida")
        return hours * 60 + minutes + (1 if seconds >= 30 else 0)
    if re.fullmatch(r"\d+", value):
        return int(value)
    raise ValueError("Se esperaban minutos enteros o HH:MM:SS")


def csv_rows(file_bytes: bytes) -> tuple[list[dict[str, Any]], int]:
    text: str
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("cp1252")
    reader = csv.DictReader(io.StringIO(text))
    fields = {re.sub(r"\s+", "", name or "").lower(): name for name in (reader.fieldnames or [])}
    required = ["idestudio", "fecha", "materia", "cantidad", "tiempo", "totalminutos"]
    if not all(name in fields for name in required):
        raise HTTPException(422, "El CSV no contiene los encabezados esperados.")
    parsed: list[dict[str, Any]] = []
    rejected = 0
    seen: set[str] = set()
    for row in reader:
        try:
            identifier = str(row[fields["idestudio"]]).strip()
            if not identifier or identifier in seen:
                raise ValueError("ID vacío o repetido")
            seen.add(identifier)
            raw_date = str(row[fields["fecha"]]).strip()
            try:
                session_date = date.fromisoformat(raw_date)
            except ValueError:
                session_date = datetime.strptime(raw_date, "%d/%m/%Y").date()
            subject = str(row[fields["materia"]]).strip()
            quantity = int(str(row[fields["cantidad"]]).strip())
            duration = parse_minutes(str(row[fields["tiempo"]]))
            total_cell = parse_minutes(str(row[fields["totalminutos"]]))
            if not subject or quantity <= 0 or duration <= 0 or quantity > 100 or duration > 600:
                raise ValueError("Datos fuera de rango")
            calculated = quantity * duration
            if total_cell != calculated:
                raise ValueError("El total no coincide con cantidad × duración")
            parsed.append({"id_estudio": identifier, "fecha": session_date, "materia": subject,
                           "cantidad": quantity, "tiempo": duration, "total_minutos": calculated})
        except (ValueError, TypeError, KeyError, OverflowError):
            rejected += 1
    return parsed, rejected


@app.post("/api/import/csv")
async def import_csv(user_id: UserId, file: UploadFile = File(...)):
    raw = await file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(413, "El archivo supera 5 MB.")
    records, rejected = csv_rows(raw)
    imported = updated = rejected = 0
    with db_connection() as connection:
        for item in records:
            previous = connection.execute("SELECT 1 FROM public.pomodoro_sessions WHERE id_estudio=%s AND user_id=%s", (item["id_estudio"], user_id)).fetchone()
            result = connection.execute(
                """INSERT INTO public.pomodoro_sessions (id_estudio,user_id,fecha,materia,cantidad,tiempo)
                   VALUES (%s,%s,%s,%s,%s,%s) ON CONFLICT(id_estudio) DO UPDATE SET
                   fecha=EXCLUDED.fecha,materia=EXCLUDED.materia,cantidad=EXCLUDED.cantidad,tiempo=EXCLUDED.tiempo,
                   updated_at=now()
                   WHERE public.pomodoro_sessions.user_id=EXCLUDED.user_id""",
                (item["id_estudio"], user_id, item["fecha"], item["materia"], item["cantidad"], item["tiempo"]),
            )
            if connection.execute("SELECT 1 FROM public.pomodoro_sessions WHERE id_estudio=%s AND user_id=%s", (item["id_estudio"], user_id)).fetchone():
                updated += int(previous is not None)
                imported += int(previous is None)
            else:
                rejected += 1
    return {"imported": imported, "updated": updated, "rejected": rejected, "received": imported + updated + rejected}


@app.get("/api/export/csv")
def export_csv(user_id: UserId):
    records = get_sessions(user_id)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["IDEstudio", "Fecha", "Materia", "Cantidad", "Tiempo (min)", "Total minutos"])
    for item in records:
        writer.writerow([item["id_estudio"], item["fecha"], item["materia"], item["cantidad"], item["tiempo"], item["total_minutos"]])
    data = ("\ufeff" + output.getvalue()).encode("utf-8")
    return StreamingResponse(io.BytesIO(data), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=pomodoro-sesiones.csv"})


@app.get("/api/export/xlsx")
def export_xlsx(user_id: UserId):
    records = get_sessions(user_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sesiones"
    sheet.append(["IDEstudio", "Fecha", "Materia", "Cantidad", "Tiempo (min)", "Total minutos"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6D5731")
        cell.alignment = Alignment(horizontal="center")
    for item in records:
        sheet.append([item["id_estudio"], date.fromisoformat(item["fecha"]), item["materia"], item["cantidad"], item["tiempo"], item["total_minutos"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEF", [40, 15, 32, 14, 17, 18]):
        sheet.column_dimensions[column].width = width
    for cell in sheet["B"][1:]:
        cell.number_format = "dd/mm/yyyy"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=pomodoro-sesiones.xlsx"})


from fastapi.staticfiles import StaticFiles
app.mount("/", StaticFiles(directory=APP_DIR, html=True), name="static")
